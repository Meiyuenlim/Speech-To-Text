from io import BytesIO

import pytest
from docx import Document


def meeting_with_transcript(client, rows):
    created = client.post("/api/meetings", json={"title": "Export review"})
    assert created.status_code == 201
    meeting_id = created.json()["id"]
    saved = client.put(
        f"/api/meetings/{meeting_id}/transcript",
        json={"segments": [
            {"sequence": index, "start_time": start, "end_time": end, "text": text}
            for index, (start, end, text) in enumerate(rows)
        ]},
    )
    assert saved.status_code == 200
    return meeting_id, saved.json()


@pytest.mark.parametrize(("rows", "expected"), [
    ([], ""),
    (
        [(0, 2, "The deadline is Friday."), (2, 4, "Please send the report.")],
        "The deadline is Friday. Please send the report.",
    ),
    (
        [(0, 2, "First topic."), (6, 8, "Next topic."), (8, 9, "谢谢大家。")],
        "First topic.\n\nNext topic. 谢谢大家。",
    ),
    (
        [(0, 4, "Corrected text.\n\n谢谢大家。"), (5, 7, " Kita akan semak."), (8, 10, "\n\nNext steps.")],
        "Corrected text.\n\n谢谢大家。 Kita akan semak.\n\nNext steps.",
    ),
    (
        [(0, 4, "Review the data"), (0, 4, "base design."), (5, 7, "Then approve it.")],
        "Review the database design. Then approve it.",
    ),
    (
        [(0, 10, "A" * 400 + "."), (10, 12, "Next paragraph.")],
        "A" * 400 + ".\n\nNext paragraph.",
    ),
], ids=["empty", "joined", "pause", "saved-paragraphs-and-unicode", "split-paste", "long-paragraph"])
def test_exports_contain_the_full_transcript_without_changing_saved_segments(client, rows, expected):
    meeting_id, saved_segments = meeting_with_transcript(client, rows)
    url = f"/api/meetings/{meeting_id}"

    text = client.get(f"{url}/export/txt")
    assert text.status_code == 200
    assert text.text.split("## Transcript\n\n", 1)[1] == expected + "\n"

    payload = client.get(f"{url}/export/json")
    assert payload.status_code == 200
    assert payload.json()["transcript"] == expected
    assert payload.json()["minutes"] is None

    word = client.get(f"{url}/export/docx")
    assert word.status_code == 200
    document = Document(BytesIO(word.content))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    transcript_start = paragraphs.index("Transcript") + 1
    assert paragraphs[transcript_start:] == (expected.split("\n\n") if expected else [])

    # Export formatting must not merge, overwrite, or remove stored timestamps.
    assert client.get(f"{url}/transcript").json() == saved_segments


def test_continuous_transcript_exports_keep_minutes_and_their_evidence(client):
    meeting_id, _ = meeting_with_transcript(client, [(0, 2, "A draft was agreed."), (2, 4, "We will review it Friday.")])
    url = f"/api/meetings/{meeting_id}"
    minutes = {
        "meeting_title": "Export review",
        "executive_summary": "The team agreed to prepare a draft.",
        "decisions_made": [{
            "text": "Prepare a draft.", "evidence_start_time": 0, "evidence_end_time": 2,
            "confidence": 0.9, "source_excerpt": "A draft was agreed.",
        }],
    }
    assert client.put(f"{url}/minutes", json={"minutes": minutes}).status_code == 200

    text = client.get(f"{url}/export/txt").text
    assert minutes["executive_summary"] in text
    assert "00:00:00" in text.split("## Transcript", 1)[0]
    assert "00:00:00" not in text.split("## Transcript", 1)[1]
    payload = client.get(f"{url}/export/json").json()
    assert payload["minutes"]["decisions_made"] == minutes["decisions_made"]
    assert payload["transcript"] == "A draft was agreed. We will review it Friday."
    document = Document(BytesIO(client.get(f"{url}/export/docx").content))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    assert "Prepare a draft. [evidence 00:00:00–00:00:02]" in paragraphs
    assert paragraphs[-1] == payload["transcript"]
