from __future__ import annotations

import re
from datetime import datetime, timezone
from io import BytesIO

from docx import Document

from app.models.meeting import Meeting
from app.schemas.meeting import MinutesDocument
from app.services.markdown import to_markdown
from app.services.transcription import format_timestamp


def full_transcript_text(meeting: Meeting) -> str:
    """Format the saved transcript like the continuous document in the review UI."""
    parts: list[str] = []
    previous = None
    paragraph_length = 0
    for segment in meeting.transcript_segments:
        separator = ""
        same_range = previous is not None and (
            previous.start_time == segment.start_time and previous.end_time == segment.end_time
        )
        if previous is not None and not same_range and not (
            re.search(r"\s$", previous.text) or re.match(r"\s", segment.text)
        ):
            paragraph_break = segment.start_time - previous.end_time >= 3 or (
                paragraph_length >= 400 and re.search(r'''[.!?。！？]["'”’）)]*$''', previous.text)
            )
            separator = "\n\n" if paragraph_break else " "
        text = separator + segment.text
        parts.append(text)
        paragraph_length = len(text.rsplit("\n\n", 1)[-1]) if "\n\n" in text else paragraph_length + len(text)
        previous = segment
    return "".join(parts)


def export_payload(meeting: Meeting, minutes: MinutesDocument | None) -> dict:
    return {
        "meeting": {
            "id": meeting.id,
            "title": meeting.title,
            "meeting_date": meeting.meeting_date.isoformat() if meeting.meeting_date else None,
            "purpose": meeting.purpose,
            "original_filename": meeting.original_filename,
            "language": meeting.language,
            "whisper_model": meeting.whisper_model,
            "ollama_model": meeting.ollama_model,
            "noise_reduction_enabled": meeting.noise_reduction_enabled,
            "status": meeting.status,
            "created_at": meeting.created_at.isoformat() if meeting.created_at else None,
        },
        "transcript": full_transcript_text(meeting),
        "minutes": minutes.model_dump(mode="json") if minutes else None,
        "exported_at": datetime.now(timezone.utc).isoformat(),
    }


def export_txt(meeting: Meeting, minutes: MinutesDocument | None) -> str:
    body = to_markdown(minutes) if minutes else f"# {meeting.title}\n\nMinutes have not been generated.\n"
    return body + "\n\n## Transcript\n\n" + full_transcript_text(meeting) + "\n"


def export_docx(meeting: Meeting, minutes: MinutesDocument | None) -> bytes:
    document = Document()
    document.add_heading(meeting.title, level=0)
    document.add_paragraph(f"Date: {meeting.meeting_date.isoformat() if meeting.meeting_date else 'Not supplied'}")
    document.add_paragraph(f"Purpose: {meeting.purpose or 'Not supplied'}")
    if minutes:
        document.add_heading("Executive summary", level=1)
        document.add_paragraph(minutes.executive_summary or "Not provided.")
        for title, items in (
            ("Main discussion points", minutes.main_discussion_points),
            ("Decisions", minutes.decisions_made),
            ("Risks and concerns", minutes.risks_and_concerns),
            ("Unresolved questions", minutes.unresolved_questions),
            ("Next steps", minutes.next_steps),
        ):
            document.add_heading(title, level=1)
            for item in items:
                document.add_paragraph(
                    f"{item.text} [evidence {format_timestamp(item.evidence_start_time)}–{format_timestamp(item.evidence_end_time)}]",
                    style="List Bullet",
                )
        document.add_heading("Action items", level=1)
        table = document.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        for cell, label in zip(table.rows[0].cells, ("Action", "Owner", "Deadline", "Status", "Evidence")):
            cell.text = label
        for item in minutes.action_items:
            cells = table.add_row().cells
            cells[0].text = item.text
            cells[1].text = item.owner or "—"
            cells[2].text = item.deadline or "—"
            cells[3].text = item.status
            cells[4].text = f"{format_timestamp(item.evidence_start_time)}–{format_timestamp(item.evidence_end_time)}"
        if minutes.generation_warnings:
            document.add_heading("Generation warnings", level=1)
            for warning in minutes.generation_warnings:
                document.add_paragraph(warning, style="List Bullet")
    document.add_heading("Transcript", level=1)
    transcript = full_transcript_text(meeting)
    if transcript:
        for paragraph in transcript.split("\n\n"):
            document.add_paragraph(paragraph)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
